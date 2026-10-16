#!/usr/bin/env python3
"""步骤2: 台站信息提取 + 到时检测"""
import sys, os, glob, argparse, yaml

CODE_ROOT = "/home/sy/codex"
for p in [CODE_ROOT, CODE_ROOT + '/t_phase_analysis']:
    if p not in sys.path: sys.path.insert(0, p)

import numpy as np
from obspy import read
from lib.arrivals import compute_theoretical_arrivals, detect_arrival_envelope, dist_az, compute_snr
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

CONFIG = yaml.safe_load(open(f"{CODE_ROOT}/config/events.yaml"))


def main():
    p = argparse.ArgumentParser()
    p.add_argument('event')
    args = p.parse_args()

    cfg = CONFIG[args.event]
    case_dir = cfg.get('case_dir', cfg['label'])
    proc_dir = f"/mnt/f/sharedate/case/{case_dir}/sacprocess"
    out_dir = f"/mnt/f/sharedate/case/{case_dir}"
    os.makedirs(out_dir, exist_ok=True)

    bhz = sorted(glob.glob(f"{proc_dir}/*.BHZ*_proc.sac"))
    print(f"事件: {cfg['label']} | {len(bhz)} BHZ文件\n")

    rows = []
    for fn in bhz:
        try:
            tr = read(fn)[0]
            sac = tr.stats.sac; stla = getattr(sac,'stla',0); stlo = getattr(sac,'stlo',0)
            if stla == 0: continue
            sta = os.path.basename(fn).split('.')[1]
            dist_km, az, baz = dist_az(cfg['lon'], cfg['lat'], stlo, stla)
            arrs = compute_theoretical_arrivals(cfg['lon'], cfg['lat'], cfg['depth_km'], stlo, stla)
            tp_th = arrs.get('T-P', 0)
            if tp_th:
                t_obs = detect_arrival_envelope(tr.data, tr.stats.delta, tp_th,
                                                 dist_km=dist_km)
                ts_th = arrs.get('T-S', 0)
                p_arrival = arrs.get('P', 0)
                t_offset = tr.stats.sac.b - tr.stats.sac.o
                snr = compute_snr(tr.data, tr.stats.delta,
                                 signal_center=ts_th, noise_center=p_arrival,
                                 signal_start=tp_th, t_offset=t_offset)
            else:
                t_obs = 0; snr = 0
            land_km = arrs.get('land_km', 0)
            ocean_km = arrs.get('ocean_km', 0)
            rows.append(dict(sta=sta, lon=stlo, lat=stla, dist=dist_km, az=az, baz=baz,
                            land_km=land_km, ocean_km=ocean_km,
                            tp_th=tp_th, tp_th_min=tp_th/60, t_obs=t_obs/60, snr=snr))
        except: continue

    rows.sort(key=lambda x: x['land_km'])

    # ── 生成 Excel ──
    wb = Workbook(); ws = wb.active; ws.title = '台站信息'

    # 事件信息行
    ws.merge_cells('A1:L1')
    info_cell = ws.cell(row=1, column=1,
        value=f"事件: {cfg['date']} Mw {cfg['mw']}  |  "
              f"{cfg['lat']}°N, {cfg['lon']}°W  |  "
              f"深度 {cfg['depth_km']} km  |  "
              f"总台站: {len(rows)}")
    info_cell.font = Font(bold=True, size=11, color='1f4e79')
    info_cell.alignment = Alignment(horizontal='left')

    # 表头（中文+英文）
    hdrs = ['台站(STA)', '经度(LON)', '纬度(LAT)', '震中距(DIST_km)', '陆地距(LAND_km)', '海洋距(OCEAN_km)',
            '方位角(AZ°)', '反方位角(BAZ°)', 'T理论(T_theory_min)', 'T实测(T_obs_min)', 'SNR(dB)']
    hdr_fill = PatternFill('solid', fgColor='4472C4')
    hdr_font = Font(bold=True, color='FFFFFF', size=10)

    for ci, h in enumerate(hdrs, 1):
        c = ws.cell(row=3, column=ci, value=h)
        c.font = hdr_font; c.fill = hdr_fill
        c.alignment = Alignment(horizontal='center', vertical='center')

    # 数据行
    for ri, r in enumerate(rows, 4):
        vals = [r['sta'], round(r['lon'],4), round(r['lat'],4),
                round(r['dist'],0), round(r['land_km'],0), round(r['ocean_km'],0),
                round(r['az'],1), round(r['baz'],1),
                round(r['tp_th_min'],2), round(r['t_obs'],2),
                round(r['snr'],2)]
        for ci, v in enumerate(vals, 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.font = Font(size=9)
            # SNR 高亮
            if ci == 11:  # SNR(dB) 列
                if isinstance(v, (int, float)) and v >= 3.0:
                    cell.fill = PatternFill('solid', fgColor='C6EFCE')
                    cell.font = Font(bold=True, size=9, color='006100')

    # 列宽
    widths = [8, 12, 12, 12, 12, 12, 10, 12, 10, 10, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64+i) if i <= 26 else 'A'+chr(64+i-26)].width = w

    # 边框
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9'))
    for row in ws.iter_rows(min_row=3, max_row=len(rows)+3, max_col=len(hdrs)):
        for cell in row:
            cell.border = thin_border

    wb.save(f"{out_dir}/stations.xlsx")

    # Save GMT-ready file
    with open(f"{out_dir}/gmt_stations.txt", 'w') as f:
        f.write("# lon lat arrival_min station\n")
        for r in rows:
            if r['t_obs'] > 0:
                f.write(f"{r['lon']:.5f} {r['lat']:.5f} {r['t_obs']:.3f} {r['sta']}\n")

    snr_ok = [r for r in rows if r['snr'] >= cfg.get('snr_threshold', 3.0)]
    print(f"总台站: {len(rows)}")
    print(f"{cfg['label']}: {len(rows)} stations | SNR>={cfg.get('snr_threshold',3.0)}dB: {len(snr_ok)}")
    print(f"输出: {out_dir}/stations.xlsx, gmt_stations.txt")

if __name__ == '__main__':
    main()
